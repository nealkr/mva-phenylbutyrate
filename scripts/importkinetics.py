#!/usr/bin/env python3
"""Fixed endpoint analysis of published paired mitochondrial import assays."""
import argparse
from collections import defaultdict
from datetime import datetime, timezone
import hashlib
import itertools
import json
import math
from pathlib import Path
import re

import numpy as np
import pandas as pd
from scipy import stats
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string


PRIMARY=['terminal_10min','incremental_auc_after_1min']
SECONDARY=['unsubtracted_auc_1_10','early_interval_change_1_2','paired_incremental_auc_ratio']


def digest(path):return hashlib.sha256(path.read_bytes()).hexdigest()


def bh(values):
    values=np.asarray(values,float);order=np.argsort(values);n=len(values)
    ranked=np.minimum.accumulate((values[order]*n/np.arange(1,n+1))[::-1])[::-1]
    out=np.empty(n);out[order]=np.minimum(ranked,1);return out


def endpoint(y,times,name):
    if name=='terminal_10min':return float(y[-1])
    if name=='incremental_auc_after_1min':return float(np.trapezoid(y-y[0],times))
    if name=='unsubtracted_auc_1_10':return float(np.trapezoid(y,times))
    if name=='early_interval_change_1_2':return float((y[1]-y[0])/(times[1]-times[0]))
    raise ValueError(name)


def summarize(wt,clone):
    wt=np.asarray(wt,float);clone=np.asarray(clone,float);difference=clone-wt;n=len(difference)
    sd=float(np.std(difference,ddof=1));se=sd/math.sqrt(n);mean=float(np.mean(difference))
    if se:
        p=float(stats.t.sf(abs(mean/se),n-1)*2);radius=float(stats.t.ppf(.975,n-1)*se)
    else:p=1.0 if mean==0 else 0.0;radius=0.0
    signs=np.asarray(list(itertools.product([-1,1],repeat=n)))
    null=np.mean(signs*difference,axis=1)
    exact=float(np.mean(np.abs(null)>=abs(mean)-1e-12))
    return {'n_pairs':n,'mean_wt':float(np.mean(wt)),'mean_clone':float(np.mean(clone)),
      'mean_paired_difference':mean,'difference_ci95_low':mean-radius,'difference_ci95_high':mean+radius,
      'paired_t_p_two_sided':p,'exact_signflip_p_two_sided':exact,'exact_two_sided_p_floor':2/(2**n),
      'paired_difference_values':json.dumps(difference.tolist()),'all_pairs_same_direction':bool(np.all(difference>0) or np.all(difference<0))}


def validate_extract(source,config):
    assert digest(source)==config['source_sha256'],'Source hash mismatch'
    book=load_workbook(source,read_only=True,data_only=False)
    tables={}
    for sheet in {p['sheet'] for p in config['pairs']}:
        assert sheet in book.sheetnames
        tables[sheet]={}
        for row in book[sheet].iter_rows(min_row=1,max_row=100,max_col=16):
            for cell in row:
                if cell.value is not None:tables[sheet][cell.coordinate]=cell.value
    book.close();records=[];validation=[]
    for pair in config['pairs']:
        table=tables[pair['sheet']];raw=pair['raw_row'];normal=pair['normalized_row'];normalizer=float(table.get('F'+str(raw)))
        assert math.isfinite(normalizer) and normalizer>0
        flag_cell=pair.get('source_quality_flag_cell');flag=table.get(flag_cell) if flag_cell else None
        if flag_cell:
            assert isinstance(flag,str) and 'pk' in flag.lower() and 'work' in flag.lower(), 'Missing literal PK source flag'
        values={}
        for arm,columns in [('wt',pair['wt_columns']),('clone',pair['clone_columns'])]:
            armvalues=[]
            for i,col in enumerate(columns):
                rawcell=f'{col}{raw}';normalcell=f'{col}{normal}';value=float(table.get(rawcell))
                assert float(table.get(f'{col}{raw-1}'))==config['times_min'][i], (pair['sheet'],col,raw-1,'time header differs')
                assert math.isfinite(value) and value>=0,(pair['sheet'],rawcell,'raw signal not numeric')
                formula=table.get(normalcell)
                if isinstance(formula,str) and formula.startswith('='):
                    assert re.search(r'\$?F\$?'+str(raw)+r'(?!\d)',formula), (normalcell,'wrong paired normalizer')
                    assert re.search(r'\$?'+col+r'\$?'+str(raw)+r'(?!\d)',formula),(normalcell,'wrong raw numerator')
                    assert '100' in formula,(normalcell,'missing percent multiplier')
                else:raise AssertionError((pair['sheet'],normalcell,'expected source normalization formula'))
                armvalues.append(float(value)*100/normalizer)
                records.append({'cargo':pair['cargo'],'background':pair['background'],'replicate':pair['replicate'],
                  'sheet':pair['sheet'],'raw_row':raw,'normalized_row':normal,'arm':arm,'time_min':config['times_min'][i],
                  'raw_signal':float(value),'paired_wt_10min_raw_signal':float(normalizer),'normalized_import_percent':armvalues[-1],
                  'source_quality_flag':str(flag) if flag else '', 'primary_eligible':not bool(flag)})
            values[arm]=armvalues
        assert abs(values['wt'][-1]-100)<1e-10
        wt_fingerprint=hashlib.sha256(json.dumps([float(table[c+str(raw)]) for c in pair['wt_columns']],separators=(',',':')).encode()).hexdigest()
        validation.append({**pair,'source_quality_flag':flag,'primary_eligible':not bool(flag),'wt_raw_trajectory_sha256':wt_fingerprint,'normalization_formulas_checked':8,'time_headers_checked':8,'time_header_row':raw-1})
    return pd.DataFrame(records),validation


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--root',type=Path,default=Path(__file__).resolve().parents[1]);args=parser.parse_args()
    source=args.root/'data/public/protein_import/source_data.xlsx';configpath=args.root/'data/public/protein_import/kinetics_config.json'
    config=json.loads(configpath.read_text());out=args.root/'results/import_kinetics';out.mkdir(parents=True,exist_ok=True)
    trajectories,validation=validate_extract(source,config)
    trajectories.to_csv(out/'paired_trajectories.tsv',sep='\t',index=False)
    times=np.array(config['times_min'],float);endpointrows=[]
    for (cargo,background,replicate),group in trajectories.groupby(['cargo','background','replicate'],sort=False):
        wt=group[group.arm=='wt'].sort_values('time_min').normalized_import_percent.to_numpy()
        clone=group[group.arm=='clone'].sort_values('time_min').normalized_import_percent.to_numpy()
        eligible=bool(group.primary_eligible.iloc[0])
        for name in PRIMARY+SECONDARY:
            if name=='paired_incremental_auc_ratio':
                w=endpoint(wt,times,'incremental_auc_after_1min');c=endpoint(clone,times,'incremental_auc_after_1min')
                assert w>0,'Nonpositive paired incremental WT AUC denominator'
                w,c=1.0,c/w
            else:w,c=endpoint(wt,times,name),endpoint(clone,times,name)
            endpointrows.append({'cargo':cargo,'background':background,'replicate':replicate,'endpoint':name,'wt':w,'clone':c,
              'paired_difference':c-w,'primary_eligible':eligible,'source_quality_flag':group.source_quality_flag.iloc[0]})
    endpoints=pd.DataFrame(endpointrows);endpoints.to_csv(out/'paired_endpoints.tsv',sep='\t',index=False)
    contrasts=[]
    for (cargo,background,name),group in endpoints.groupby(['cargo','background','endpoint'],sort=False):
        clean=group[group.primary_eligible]
        family=('primary_su9' if cargo=='su9_DHFR' else 'secondary_SOD2') if name in PRIMARY else 'endpoint_sensitivities'
        result={'cargo':cargo,'background':background,'endpoint':name,'assay_filter':'valid_PK','family':family,**summarize(clean.wt,clean.clone)}
        contrasts.append(result)
        if len(clean)<len(group) and name in PRIMARY:
            contrasts.append({'cargo':cargo,'background':background,'endpoint':name,'assay_filter':'all_source_including_failed_PK','family':'source_quality_sensitivity',**summarize(group.wt,group.clone)})
    contrasts=pd.DataFrame(contrasts)
    for family,indices in contrasts.groupby('family').groups.items():
        contrasts.loc[indices,'paired_t_q_bh']=bh(contrasts.loc[indices,'paired_t_p_two_sided'])
        contrasts.loc[indices,'signflip_q_bh']=bh(contrasts.loc[indices,'exact_signflip_p_two_sided'])
    contrasts.to_csv(out/'endpoint_contrasts.tsv',sep='\t',index=False)
    reuse=[]
    groups=defaultdict(list)
    for item in validation:groups[(item['cargo'],item['wt_raw_trajectory_sha256'])].append(item)
    for (cargo,fingerprint),items in groups.items():
        if len(items)>1:reuse.append({'cargo':cargo,'wt_raw_trajectory_sha256':fingerprint,'source_pairs':[{'background':v['background'],'replicate':v['replicate'],'raw_row':v['raw_row']} for v in items]})
    summary={'generated_utc':datetime.now(timezone.utc).isoformat(),'source_sha256':digest(source),'config_sha256':digest(configpath),
      'source_url':config['source_url'],'timepoints_min':config['times_min'],'measured_zero':False,
      'normalization':config['normalization'],'units':config['normalized_units'],
      'raw_units':'Source band-intensity units (uncalibrated)',
      'validation':validation,'shared_raw_wt_trajectories':reuse,'source_pair_count':len(validation),
      'primary_eligible_pair_count':sum(v['primary_eligible'] for v in validation),
      'primary_quality_exclusion_count':sum(not v['primary_eligible'] for v in validation),
      'family_sizes':dict(contrasts.groupby('family').size().astype(int)),
      'primary_and_SOD2_contrasts':json.loads(contrasts[contrasts.family.isin(['primary_su9','secondary_SOD2'])].to_json(orient='records')),
      'interpretation':'Relative imported signal and incremental accumulation, not absolute throughput. Source cargo/background comparisons are distinct; shared WT trajectories are not pooled independent controls. No drug arm.',
      'inference_notes':'Paired t intervals assume approximately normal paired differences; exact sign flips use a paired sign-symmetry reference distribution. Repeated timepoints are never biological n.',
      'burden_backgrounds':config['background_extra_chromosome_copies'],
      'burden_interpretation':'Only three different chromosome backgrounds; report heterogeneity without fitting a general dosage-burden response.',
      'reproduction_scope':'Fixed endpoint reanalysis of previously published mitochondrial import impairment; not independent biological confirmation'}
    for key in ['family_sizes']:summary[key]={k:int(v) for k,v in summary[key].items()}
    (out/'summary.json').write_text(json.dumps(summary,indent=2,allow_nan=False)+'\n')
    print(json.dumps({'output':str(out),'source_pairs':len(validation),'valid_pairs':summary['primary_eligible_pair_count'],'contrasts':len(contrasts),'families':summary['family_sizes']}))


if __name__=='__main__':main()
