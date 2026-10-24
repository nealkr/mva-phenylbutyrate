"""Bounded reanalysis of downloaded NAC functional observations.

Requires numpy, scipy and openpyxl. Input files remain unchanged. Biological
units are pooled larval chambers (Rodriguez) or independent experiments (Wang).
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import itertools
import json
import math
from pathlib import Path

import numpy as np
import openpyxl


def values(rows, field):
    return np.array([float(r[field]) for r in rows if r.get(field, "") != ""], dtype=float)


def interval(a, b):
    """Unpaired Welch estimate: mean(a)-mean(b), CI and approximate p."""
    from scipy import stats

    a, b = np.asarray(a, float), np.asarray(b, float)
    v = [a.var(ddof=1) / len(a), b.var(ddof=1) / len(b)]
    variance = sum(v)
    se = math.sqrt(variance)
    df = variance ** 2 / sum(x ** 2 / (n - 1) for x, n in zip(v, [len(a), len(b)]))
    effect = float(a.mean() - b.mean())
    half = float(stats.t.ppf(0.975, df) * se)
    p = float(2 * stats.t.sf(abs(effect / se), df))
    return {"mean_treated": float(a.mean()), "mean_reference": float(b.mean()),
            "difference": effect, "ci95": [effect - half, effect + half],
            "welch_df": df, "welch_p": p}


def permutation(a, b):
    """Enumerate all between-arm label allocations; absolute mean difference."""
    z = np.r_[a, b]
    n = len(a)
    observed = abs(float(np.mean(a) - np.mean(b)))
    total, extreme = 0, 0
    for ix in itertools.combinations(range(len(z)), n):
        left = z[list(ix)].sum()
        effect = left / n - (z.sum() - left) / (len(z) - n)
        total += 1
        extreme += abs(effect) >= observed - 1e-12
    return {"exact_permutation_p": extreme / total, "label_allocations": total,
            "permutation_assumption": "exchangeable unpaired biological units between arms"}


def contrast(dataset, endpoint, stratum, treated, reference, unit, units, label):
    a = np.array([r[1] for r in treated], float)
    b = np.array([r[1] for r in reference], float)
    out = {"dataset": dataset, "endpoint": endpoint, "stratum": stratum,
           "comparison": label, "n_treated": len(a), "n_reference": len(b),
           "unit": unit, "measurement_units": units,
           "treated_ids": [r[0] for r in treated], "reference_ids": [r[0] for r in reference],
           "treated_observations": a.tolist(), "reference_observations": b.tolist()}
    out.update(interval(a, b))
    out.update(permutation(a, b))
    out["multiplicity_p"] = out["exact_permutation_p"]
    return out


def difference_in_differences(dataset, endpoint, genotype, gt, gc, wt, wc, units):
    """Four-group Welch contrast; no across-genotype pooling or endpoint pairing."""
    from scipy import stats

    groups = [np.array(x, float) for x in [gt, gc, wt, wc]]
    weights = [1, -1, -1, 1]
    effect = float(sum(w * a.mean() for w, a in zip(weights, groups)))
    terms = [a.var(ddof=1) / len(a) for a in groups]
    variance = sum(terms)
    se = math.sqrt(variance)
    df = variance ** 2 / sum(v ** 2 / (len(a) - 1) for v, a in zip(terms, groups))
    half = float(stats.t.ppf(0.975, df) * se)
    p = float(2 * stats.t.sf(abs(effect / se), df))
    return {"dataset": dataset, "endpoint": endpoint, "stratum": genotype,
            "comparison": "(NAC-Control) in " + genotype + " minus (NAC-Control) in WT",
            "difference": effect, "ci95": [effect - half, effect + half],
            "measurement_units": units, "welch_df": df, "welch_p": p,
            "multiplicity_p": p, "method": "four-group unpaired Welch contrast",
            "group_n": [len(a) for a in groups],
            "permutation_p": None,
            "permutation_reason": "simple within-genotype label permutations do not test equal nonzero treatment effects"}


def bh(rows):
    order = sorted(range(len(rows)), key=lambda i: rows[i]["multiplicity_p"])
    previous = 1.0
    for rank in range(len(order), 0, -1):
        i = order[rank - 1]
        previous = min(previous, rows[i]["multiplicity_p"] * len(order) / rank)
        rows[i]["bh_q"] = previous
        rows[i]["bh_family_n"] = len(order)


def rodriguez(root, config, validation_only=False):
    results = []
    audits = []
    for spec in config["endpoints"]:
        with (root / spec["file"]).open(newline="", encoding="utf-8-sig") as handle:
            rows = list(csv.DictReader(handle))
        selected = {}
        for genotype in ["WT", "COX", "BAR"]:
            selected[genotype] = {}
            for arm in ["Control", "NAC"]:
                selected[genotype][arm] = [(r["code"], float(r[spec["column"]]))
                    for r in rows if r["mito"] == genotype and r["treat"] == arm and r[spec["column"]] != ""]
                expected = spec["expected_n"][genotype][arm]
                assert len(selected[genotype][arm]) == expected, (spec["column"], genotype, arm)
                assert len({r[0] for r in selected[genotype][arm]}) == expected
            if not validation_only:
                results.append(contrast("Rodriguez2021", spec["column"], genotype,
                    selected[genotype]["NAC"], selected[genotype]["Control"],
                    "one chamber containing two pooled larvae", spec["units"], "NAC-Control"))
        if not validation_only:
            for genotype in ["COX", "BAR"]:
                results.append(difference_in_differences("Rodriguez2021", spec["column"], genotype,
                    [x[1] for x in selected[genotype]["NAC"]], [x[1] for x in selected[genotype]["Control"]],
                    [x[1] for x in selected["WT"]["NAC"]], [x[1] for x in selected["WT"]["Control"]], spec["units"]))
        all_arm_n = {g + "|" + a: sum(r["mito"] == g and r["treat"] == a for r in rows)
                     for g in ["WT", "COX", "BAR"] for a in ["Control", "NAC", "Protein"]}
        audits.append({"file": spec["file"], "rows": len(rows), "selected_n": spec["expected_n"],
                       "all_deposited_arm_counts": all_arm_n, "qc_exclusions": [],
                       "scope": "All deposited NAC and Control observations used; Protein comparator retained in source, outside authorized NAC-Control contrasts."})
    if not validation_only:
        bh(results)
    return results, audits


def wang(root, config, validation_only=False):
    results, audits = [], []
    for spec in config["endpoints"]:
        sheet = openpyxl.load_workbook(root / spec["file"], data_only=True)[spec["sheet"]]
        columns = spec["columns"]
        labels = [sheet.cell(spec["header_row"], col).value for col in columns]
        assert labels == spec["arm_labels"], (spec["endpoint"], labels)
        arms = {}
        for col, arm in zip(columns, labels):
            arms[arm] = [(str(sheet.cell(row, 2).value), float(sheet.cell(row, col).value))
                         for row in spec["data_rows"]]
        assert all([x[0] for x in a] == spec["experiment_labels"] for a in arms.values())
        comparisons = [("shNDUFA3+NAC", "shNDUFA3"), ("Control+NAC", "Control")]
        if "Control+ATP" in arms:
            comparisons += [("shNDUFA3+ATP", "shNDUFA3"), ("Control+ATP", "Control")]
        if not validation_only:
            for treated, reference in comparisons:
                results.append(contrast("Wang2026", spec["endpoint"], reference, arms[treated], arms[reference],
                    "independent experiment, not individual cell", "percentage points", treated + " minus " + reference))
        audits.append({"file": spec["file"], "endpoint": spec["endpoint"], "labels": labels,
                       "n_per_arm": len(spec["data_rows"]), "row_pairing_verified": False})
    if not validation_only:
        bh(results)
    return results, audits


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data-dir", type=Path, default=Path(__file__).resolve().parent)
    parser.add_argument("--config", type=Path)
    parser.add_argument("--dataset", choices=["all", "rodriguez", "wang"], default="all")
    parser.add_argument("--output", type=Path, default=Path("nac_functional_results.json"))
    parser.add_argument("--validate-only", action="store_true")
    args = parser.parse_args()
    cfg = json.loads((args.config or args.data_dir / "nac_functional_config.json").read_text(encoding="utf-8"))
    for file, expected in cfg["input_sha256"].items():
        actual = hashlib.sha256((args.data_dir / file).read_bytes()).hexdigest()
        assert actual == expected, (file, actual, expected)
    out = {"config": cfg, "datasets": {}}
    for name, method in [("rodriguez", rodriguez), ("wang", wang)]:
        if args.dataset in ["all", name]:
            rows, audit = method(args.data_dir, cfg[name], args.validate_only)
            out["datasets"][name] = {"audit": audit, "contrasts": [] if args.validate_only else rows}
    if args.validate_only:
        print(json.dumps({"validated": True, "audits": {k: v["audit"] for k, v in out["datasets"].items()}}, indent=2))
    else:
        args.output.write_text(json.dumps(out, indent=2, allow_nan=False) + "\n", encoding="utf-8")
        print(str(args.output))


if __name__ == "__main__":
    main()
